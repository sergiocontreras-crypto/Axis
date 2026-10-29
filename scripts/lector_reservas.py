# -*- coding: utf-8 -*-
"""
================================================================================
LECTOR DE RESERVAS PARA LA PESTAÑA ENCAJE DEL DASHBOARD AXIS
================================================================================

PROPOSITO
---------
Provee al dashboard (app_dash.py) los SALDOS de reservas y el INTERES a la
reserva por ramo, ya homologados al diccionario de portafolios, seleccionando
el mes correcto segun la fecha del sidebar y proyectando hacia adelante cuando
el mes objetivo cae fuera del rango disponible en el insumo.

FUENTE
------
    Plantillas Input/Inputs_Reservas e interes a la reserva actualizado.xlsx
      - hoja 'Saldo Reservas'      : col C = Dic-2025 (base), D..O = Ene..Dic 2026
      - hoja 'Interés a la Reserva': col C..N = Ene..Dic 2026 (valores en millones)

REGLAS DE NEGOCIO (confirmadas con Gerencia Financiera)
-------------------------------------------------------
1. MES CAIDO: la reserva a usar depende del dia de la fecha de INICIO del
   sidebar:
       dia >= 20  -> mes anterior            (M-1)
       dia <  20  -> dos meses atras         (M-2)
   Ej: inicio 20-jul -> junio ; inicio 19-jul -> mayo.

2. SALDO DIRECTO: el valor del mes objetivo se toma TAL CUAL de la columna del
   insumo (no se suman variaciones).

3. PROYECCION POR CRECIMIENTO (cuando el mes objetivo esta fuera del rango del
   insumo): se calcula, POR CADA RAMO, la tasa de crecimiento mensual promedio
   GEOMETRICA (CAGR mensual) observada sobre la serie disponible, y se proyecta
   mes a mes con esa tasa propia hasta el mes objetivo. La FECHA FIN del sidebar
   limita hasta que mes se permite proyectar.

4. Aplica a AMBAS hojas: saldos e interes. Cada ramo usa su propia tasa.

Autor: Gerencia Financiera - AXIS
"""

import sys
from pathlib import Path
from datetime import date

import openpyxl
import pandas as pd

# ------------------------------------------------------------------------------
# RUTAS Y REUTILIZACION DE LOGICA DE HOMOLOGACION
# ------------------------------------------------------------------------------
SCRIPT_DIR = Path(__file__).resolve().parent            # <raiz>/scripts
PROYECTO_DIR = SCRIPT_DIR.parent                        # raiz del proyecto
RUTA_INSUMO = PROYECTO_DIR / "Plantillas Input" / "Inputs_Reservas e interes a la reserva actualizado.xlsx"

sys.path.insert(0, str(SCRIPT_DIR))
# leer_saldos ya homologa y agrupa por (compania, ramo) devolviendo
# { (compania, ramo): [13 saldos]  # idx 0 = Dic25, 1..12 = Ene..Dic 2026 }
# La homologacion proviene de la Fuente_Unica_De_Verdad (tablas SQLite) via el
# Lector_Mapeo; ya NO se importa el diccionario hardcodeado HOMOLOGACION_RESERVAS_A_DICCIONARIO.
from generar_inputs_reservas import (
    leer_saldos, _num, _derivar_compania_interes, _normalizar_ramo_local,
)
from compras_ventas_reservas import cargar_mapeo_interes

# ------------------------------------------------------------------------------
# EJE TEMPORAL DEL INSUMO
# ------------------------------------------------------------------------------
# Saldos: la serie tiene 13 puntos mensuales cuyo primer punto es Dic-2025.
SALDO_ANCHOR = (2025, 12)      # (anio, mes) del primer punto de la serie de saldos
# Interes: la serie tiene 12 puntos mensuales cuyo primer punto es Ene-2026.
INTERES_ANCHOR = (2026, 1)

# Estructura de columnas de la hoja de interés (el mapeo segmento→ramo vive ahora en la
# Fuente_Unica_De_Verdad y se consulta vía cargar_mapeo_interes()).
try:
    from generar_inputs_reservas import COL_SEGMENTO_INTERES, COLS_MESES_INTERES
except Exception:  # pragma: no cover - fallback defensivo
    COL_SEGMENTO_INTERES = 1
    COLS_MESES_INTERES = list(range(2, 14))


# ==============================================================================
# UTILIDADES DE FECHA
# ==============================================================================
def _a_indice_mes(anio: int, mes: int, anchor: tuple) -> int:
    """Convierte (anio, mes) al indice 0-based dentro de la serie anclada en `anchor`."""
    a0, m0 = anchor
    return (anio - a0) * 12 + (mes - m0)


def _sumar_meses(anio: int, mes: int, delta: int) -> tuple:
    """Devuelve (anio, mes) resultante de sumar `delta` meses (delta puede ser negativo)."""
    total = (anio * 12 + (mes - 1)) + delta
    return total // 12, (total % 12) + 1


def resolver_mes_objetivo(fecha_inicio) -> tuple:
    """
    Aplica la regla de MES CAIDO a la fecha de inicio del sidebar.
        dia >= 20 -> (M-1) ; dia < 20 -> (M-2)
    Devuelve (anio, mes) del mes de reserva a usar.
    """
    f = pd.to_datetime(fecha_inicio)
    delta = -1 if f.day >= 20 else -2
    return _sumar_meses(f.year, f.month, delta)


def _mes_de_fecha(fecha) -> tuple:
    """Devuelve (anio, mes) de una fecha cualquiera."""
    f = pd.to_datetime(fecha)
    return f.year, f.month


# ==============================================================================
# PROYECCION GEOMETRICA POR RAMO
# ==============================================================================
def _cagr_mensual(serie) -> float:
    """
    Tasa de crecimiento mensual GEOMETRICA sobre los puntos validos y positivos
    de `serie` (lista con posibles None).
        cagr = (ultimo / primero) ** (1 / n_pasos) - 1
    Devuelve 0.0 si no hay suficientes puntos positivos para estimarla.
    """
    vals = [v for v in serie if v is not None and v > 0]
    if len(vals) < 2:
        return 0.0
    primero, ultimo = vals[0], vals[-1]
    n_pasos = len(vals) - 1
    if primero <= 0 or ultimo <= 0 or n_pasos <= 0:
        return 0.0
    return (ultimo / primero) ** (1.0 / n_pasos) - 1.0


def _valor_en_indice(serie, idx_objetivo: int, tope_idx: int):
    """
    Devuelve el valor de `serie` en el indice mensual `idx_objetivo`.
      - Si idx_objetivo cae dentro de la serie y tiene dato -> valor directo.
      - Si cae fuera (o el dato es None), se proyecta con CAGR propio de la serie
        partiendo del ultimo (o primer) punto valido disponible.
      - `tope_idx` limita hasta que indice se permite proyectar (fecha fin del
        sidebar). Si idx_objetivo lo supera, se recorta a `tope_idx`.
    Devuelve None si no hay ningun punto valido para anclar la proyeccion.
    """
    n = len(serie)
    # Recortar el objetivo al tope permitido por la fecha fin.
    idx = min(idx_objetivo, tope_idx) if tope_idx is not None else idx_objetivo

    # Caso directo: dentro de rango y con dato.
    if 0 <= idx < n and serie[idx] is not None:
        return float(serie[idx])

    # Necesitamos proyectar. Ubicar puntos validos.
    indices_validos = [i for i, v in enumerate(serie) if v is not None and v > 0]
    if not indices_validos:
        # sin datos positivos; intentar cualquier no-None
        indices_no_none = [i for i, v in enumerate(serie) if v is not None]
        if not indices_no_none:
            return None
        # sin base para crecer: devolver el punto no-nulo mas cercano
        i_cercano = min(indices_no_none, key=lambda i: abs(i - idx))
        return float(serie[i_cercano])

    cagr = _cagr_mensual(serie)

    if idx > indices_validos[-1]:
        # Proyeccion hacia ADELANTE desde el ultimo punto valido.
        i_base = indices_validos[-1]
        base = float(serie[i_base])
        pasos = idx - i_base
        return base * ((1.0 + cagr) ** pasos)
    elif idx < indices_validos[0]:
        # Proyeccion hacia ATRAS desde el primer punto valido.
        i_base = indices_validos[0]
        base = float(serie[i_base])
        pasos = i_base - idx
        return base / ((1.0 + cagr) ** pasos) if (1.0 + cagr) != 0 else base
    else:
        # Hueco interno (None entre datos): interpolar con crecimiento desde el
        # ultimo punto valido anterior.
        anteriores = [i for i in indices_validos if i < idx]
        i_base = anteriores[-1]
        base = float(serie[i_base])
        pasos = idx - i_base
        return base * ((1.0 + cagr) ** pasos)


# ==============================================================================
# CARGA DE SALDOS DE RESERVAS (con mes caido + proyeccion)
# ==============================================================================
def _leer_filas_saldo():
    """Carga las filas de la hoja 'Saldo Reservas' indexadas por numero de fila."""
    wb = openpyxl.load_workbook(str(RUTA_INSUMO), data_only=True, read_only=True)
    ws = wb["Saldo Reservas"]
    filas = {}
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=60, values_only=True), start=1):
        filas[i] = row
    wb.close()
    return filas


def _cargar_reservas_para_mes(anio_obj, mes_obj, fecha_fin=None) -> pd.DataFrame:
    """
    Nucleo comun: devuelve los saldos de reservas por ramo para el mes de
    RESERVA (anio_obj, mes_obj) ya resuelto, proyectando con CAGR por ramo si
    el mes cae fuera del rango del insumo. `fecha_fin` limita hasta donde
    proyectar.

    Salida: DataFrame[Ramo, Saldo_Actual, Compania].
    """
    if not RUTA_INSUMO.exists():
        return pd.DataFrame(columns=["Ramo", "Saldo_Actual", "Compania"])

    filas = _leer_filas_saldo()
    saldos = leer_saldos(filas)   # { (compania, ramo): [13 saldos] } homologado

    idx_obj = _a_indice_mes(anio_obj, mes_obj, SALDO_ANCHOR)

    # Tope de proyeccion segun fecha fin (indice del mes de la fecha fin).
    tope_idx = None
    if fecha_fin is not None:
        a_fin, m_fin = _mes_de_fecha(fecha_fin)
        tope_idx = _a_indice_mes(a_fin, m_fin, SALDO_ANCHOR)

    registros = []
    for (compania, ramo), serie in saldos.items():
        valor = _valor_en_indice(serie, idx_obj, tope_idx)
        if valor is None:
            continue
        registros.append({
            "Ramo": ramo,
            "Saldo_Actual": float(valor),
            "Compania": compania,
        })

    df = pd.DataFrame(registros, columns=["Ramo", "Saldo_Actual", "Compania"])
    if df.empty:
        return df
    # Agrupar por si la homologacion colapso ramos.
    df = df.groupby(["Ramo", "Compania"], as_index=False).agg({"Saldo_Actual": "sum"})
    return df


def cargar_reservas(fecha_inicio, fecha_fin=None) -> pd.DataFrame:
    """
    Modo ACTUAL: saldos de reservas por ramo para el mes objetivo derivado de
    `fecha_inicio` con la regla de MES CAIDO (día >= 20 → M-1, día < 20 → M-2).

    Salida: DataFrame[Ramo, Saldo_Actual, Compania]  (compatible con el dashboard).
    """
    anio_obj, mes_obj = resolver_mes_objetivo(fecha_inicio)
    return _cargar_reservas_para_mes(anio_obj, mes_obj, fecha_fin)


def cargar_reservas_mes_portafolio(anio_portafolio, mes_portafolio, fecha_fin=None) -> pd.DataFrame:
    """
    Modo PROYECCION (dropdown): dado el mes de PORTAFOLIO seleccionado, la
    reserva a usar es el mes ANTERIOR (M-1, mes caído de proyección). Ej:
    portafolio = agosto → reserva de julio.

    Salida: DataFrame[Ramo, Saldo_Actual, Compania].
    """
    anio_res, mes_res = _sumar_meses(anio_portafolio, mes_portafolio, -1)
    return _cargar_reservas_para_mes(anio_res, mes_res, fecha_fin)


# ==============================================================================
# CARGA DE INTERES A LA RESERVA (con mes caido + proyeccion)
# ==============================================================================
def _leer_series_interes():
    """
    Lee la hoja 'Interés a la Reserva' y devuelve:
        { (compania, ramo homologado): [12 valores]  # idx 0 = Ene-2026 }
    usando el `MapeoDeReservas` de interés (Fuente_Unica_De_Verdad). Ignora segmentos
    sin entrada explícita en el mapeo ('Bolivar', 'TOTAL COMPAÑIAS', etc.). La compañía
    se deriva del ramo destino.
    """
    mapeo_interes = cargar_mapeo_interes(str(RUTA_INSUMO))
    wb = openpyxl.load_workbook(str(RUTA_INSUMO), data_only=True, read_only=True)
    ws = wb["Interés a la Reserva"]
    series = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        seg = row[COL_SEGMENTO_INTERES] if len(row) > COL_SEGMENTO_INTERES else None
        if seg is None:
            continue
        seg_str = str(seg).strip()
        clave = _normalizar_ramo_local(seg_str)
        ramo = mapeo_interes.entradas.get(clave)
        if not ramo:  # None (ausente) o "" (ignorar)
            continue
        compania = _derivar_compania_interes(ramo)
        vals = [_num(row[j] if len(row) > j else None) for j in COLS_MESES_INTERES]
        series[(compania, ramo)] = vals
    wb.close()
    return series


def cargar_interes(fecha_inicio, fecha_fin=None) -> pd.DataFrame:
    """
    Devuelve el interes a la reserva por ramo para el mes objetivo (regla mes
    caido sobre `fecha_inicio`), proyectando con CAGR por ramo si el mes cae
    fuera del rango del insumo. `fecha_fin` limita hasta donde proyectar.

    Salida: DataFrame[Ramo, Interes, Compania]  (valores en millones de COP).
    """
    if not RUTA_INSUMO.exists():
        return pd.DataFrame(columns=["Ramo", "Interes", "Compania"])

    series = _leer_series_interes()

    anio_obj, mes_obj = resolver_mes_objetivo(fecha_inicio)
    idx_obj = _a_indice_mes(anio_obj, mes_obj, INTERES_ANCHOR)

    tope_idx = None
    if fecha_fin is not None:
        a_fin, m_fin = _mes_de_fecha(fecha_fin)
        tope_idx = _a_indice_mes(a_fin, m_fin, INTERES_ANCHOR)

    registros = []
    for (compania, ramo), serie in series.items():
        valor = _valor_en_indice(serie, idx_obj, tope_idx)
        if valor is None:
            continue
        registros.append({
            "Ramo": ramo,
            "Interes": float(valor),
            "Compania": compania,
        })

    df = pd.DataFrame(registros, columns=["Ramo", "Interes", "Compania"])
    if df.empty:
        return df
    df = df.groupby(["Ramo", "Compania"], as_index=False).agg({"Interes": "sum"})
    return df


if __name__ == "__main__":
    # Prueba rapida con la fecha por defecto del dashboard.
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    for f_ini in ["2026-07-31", "2026-07-19", "2026-01-15", "2026-12-31"]:
        anio, mes = resolver_mes_objetivo(f_ini)
        print(f"\nfecha_inicio={f_ini} -> mes objetivo reserva: {anio}-{mes:02d}")
        dfr = cargar_reservas(f_ini, "2026-12-31")
        print(f"  reservas: {len(dfr)} ramos, total={dfr['Saldo_Actual'].sum():,.0f}")
        dfi = cargar_interes(f_ini, "2026-12-31")
        print(f"  interes : {len(dfi)} ramos, total={dfi['Interes'].sum():,.2f}")
